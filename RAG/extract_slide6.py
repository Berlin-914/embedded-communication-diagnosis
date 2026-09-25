from pathlib import Path
from pptx import Presentation

PROJECT_ROOT = Path(__file__).resolve().parent.parent

PPT_PATH = PROJECT_ROOT / "DATA" / "Unit 5 Pointing and Acquisition.pptx"
OUTPUT_DIR = PROJECT_ROOT / "DATA" / "slide6"

OUTPUT_DIR.mkdir(exist_ok=True)

prs = Presentation(PPT_PATH)

slide = prs.slides[5]   # Slide 6

print("Slide 6 found.")
print(f"Number of shapes: {len(slide.shapes)}")

for i, shape in enumerate(slide.shapes, start=1):

    print(
        f"{i}: type={shape.shape_type}, "
        f"name={shape.name}"
    )

    if shape.shape_type == 13:  # Picture

        image_path = OUTPUT_DIR / f"slide6_image_{i}.png"

        with open(image_path, "wb") as f:
            f.write(shape.image.blob)

        print(f"Image saved: {image_path}")

print("\nDone.")